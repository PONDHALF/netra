from __future__ import annotations

import csv
import io
import re
import shutil
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Query as SAQuery
from sqlalchemy.orm import Session

from ..config import CORRECTIONS_DIR, DATA_DIR
from ..db import get_db
from ..models import Event, now
from ..schemas import EventOut, EventPage, EventPatch

router = APIRouter()

VEHICLE_TYPES = {"car", "motorcycle", "bus", "truck"}
TYPE_TH = {"car": "รถยนต์", "motorcycle": "จักรยานยนต์", "bus": "รถบัส", "truck": "รถบรรทุก"}


class Filters:
    def __init__(
        self,
        q: str | None = Query(None, description="ค้นหาเลขทะเบียน/จังหวัด (บางส่วนได้ เช่น 1234)"),
        type: str | None = Query(None, description="car,motorcycle,bus,truck (คั่นด้วย ,)"),
        from_: datetime | None = Query(None, alias="from"),
        to: datetime | None = None,
        job_id: int | None = None,
        source_id: int | None = None,
        min_conf: float | None = Query(None, ge=0, le=1),
        offset_from: float | None = None,
        offset_to: float | None = None,
        plate: str | None = Query(None, description="read | unread | corrected"),
    ):
        self.q, self.type, self.from_, self.to = q, type, from_, to
        self.job_id, self.source_id, self.min_conf = job_id, source_id, min_conf
        self.offset_from, self.offset_to, self.plate = offset_from, offset_to, plate

    def apply(self, query: SAQuery) -> SAQuery:
        if self.q and self.q.strip():
            needle = re.sub(r"\s+", "", self.q.strip())
            query = query.filter(or_(func.replace(Event.plate_text, " ", "").contains(needle),
                                     Event.plate_province.contains(self.q.strip())))
        if self.type:
            types = [t for t in self.type.split(",") if t in VEHICLE_TYPES]
            if types:
                query = query.filter(Event.vehicle_type.in_(types))
        if self.from_:
            query = query.filter(Event.ts >= self.from_.replace(tzinfo=None))
        if self.to:
            query = query.filter(Event.ts <= self.to.replace(tzinfo=None))
        if self.job_id is not None:
            query = query.filter(Event.job_id == self.job_id)
        if self.source_id is not None:
            query = query.filter(Event.source_id == self.source_id)
        if self.min_conf is not None:
            query = query.filter(or_(Event.plate_conf >= self.min_conf, Event.is_corrected.is_(True)))
        if self.offset_from is not None:
            query = query.filter(Event.video_offset_sec >= self.offset_from)
        if self.offset_to is not None:
            query = query.filter(Event.video_offset_sec <= self.offset_to)
        if self.plate == "read":
            query = query.filter(Event.plate_text.isnot(None))
        elif self.plate == "unread":
            query = query.filter(Event.plate_text.is_(None))
        elif self.plate == "corrected":
            query = query.filter(Event.is_corrected.is_(True))
        return query


def _order(query: SAQuery, sort: str) -> SAQuery:
    if sort == "offset":
        return query.order_by(Event.video_offset_sec, Event.id)
    if sort == "conf":
        return query.order_by(Event.plate_conf.desc(), Event.id)
    return query.order_by(Event.ts.desc(), Event.id.desc())


def _out(e: Event) -> EventOut:
    o = EventOut.model_validate(e)
    o.source_name = e.source.name if e.source else None
    return o


@router.get("/api/events", response_model=EventPage)
def list_events(
    f: Filters = Depends(),
    sort: str = Query("ts", pattern="^(ts|offset|conf)$"),
    limit: int = Query(200, ge=1, le=5000),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
):
    query = f.apply(db.query(Event))
    total = query.count()
    items = _order(query, sort).offset(offset).limit(limit).all()
    return EventPage(items=[_out(e) for e in items], total=total)


@router.get("/api/events/export")
def export_events(
    f: Filters = Depends(),
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    sort: str = Query("ts", pattern="^(ts|offset|conf)$"),
    db: Session = Depends(get_db),
):
    events = _order(f.apply(db.query(Event)), sort).limit(50000).all()
    headers = ["ID", "วันเวลา", "เวลาในวิดีโอ", "กล้อง/สถานที่", "ประเภทรถ", "เลขทะเบียน", "จังหวัด",
               "ความมั่นใจ (%)", "แก้ไขด้วยมือ", "รูปรถ", "รูปป้าย"]

    def row(e: Event) -> list:
        off = e.video_offset_sec or 0
        return [e.id, e.ts.strftime("%Y-%m-%d %H:%M:%S"), f"{int(off // 3600):02d}:{int(off % 3600 // 60):02d}:{int(off % 60):02d}",
                e.source.name if e.source else "", TYPE_TH.get(e.vehicle_type, e.vehicle_type), e.plate_text or "",
                e.plate_province or "", round(e.plate_conf * 100, 1), "ใช่" if e.is_corrected else "",
                e.car_img or "", e.plate_img or ""]

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    if format == "csv":
        buf = io.StringIO()
        buf.write("﻿")  # BOM ให้ Excel เปิดภาษาไทยถูก
        w = csv.writer(buf)
        w.writerow(headers)
        for e in events:
            w.writerow(row(e))
        data = io.BytesIO(buf.getvalue().encode("utf-8"))
        return StreamingResponse(data, media_type="text/csv; charset=utf-8",
                                 headers={"Content-Disposition": f'attachment; filename="netra-{stamp}.csv"'})

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "NETRA"
    ws.append(headers + ["ภาพป้าย"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E293B")
        c.alignment = Alignment(horizontal="center", vertical="center")
    widths = [7, 20, 12, 22, 12, 14, 18, 14, 12, 40, 40, 26]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    embed = len(events) <= 2000
    for r, e in enumerate(events, start=2):
        ws.append(row(e))
        if embed and e.plate_img and (DATA_DIR / e.plate_img).exists():
            img = XLImage(str(DATA_DIR / e.plate_img))
            scale = 40 / max(img.height, 1)
            img.height, img.width = 40, int(img.width * scale)
            ws.add_image(img, f"L{r}")
            ws.row_dimensions[r].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="netra-{stamp}.xlsx"'})


@router.get("/api/events/{event_id}", response_model=EventOut)
def get_event(event_id: int, db: Session = Depends(get_db)):
    e = db.get(Event, event_id)
    if e is None:
        raise HTTPException(404, "ไม่พบรายการ")
    return _out(e)


@router.patch("/api/events/{event_id}", response_model=EventOut)
def patch_event(event_id: int, body: EventPatch, db: Session = Depends(get_db)):
    e = db.get(Event, event_id)
    if e is None:
        raise HTTPException(404, "ไม่พบรายการ")
    text = (re.sub(r"\s+", " ", body.plate_text.strip()) or None) if body.plate_text is not None else e.plate_text
    prov = (body.plate_province.strip() or None) if body.plate_province is not None else e.plate_province
    changed = (text, prov) != (e.plate_text, e.plate_province)
    if changed:
        if not e.is_corrected:  # เก็บค่าที่ AI อ่านได้ไว้ครั้งแรกเท่านั้น
            e.original_plate_text, e.original_plate_province = e.plate_text, e.plate_province
        e.plate_text, e.plate_province = text, prov
    if body.vehicle_type is not None and body.vehicle_type in VEHICLE_TYPES:
        e.vehicle_type = body.vehicle_type
    if changed:
        e.is_corrected, e.corrected_at, e.plate_conf = True, now(), 1.0
    db.commit()
    if changed:
        _export_corrections(db)
    return _out(e)


def _export_corrections(db: Session) -> None:
    """เก็บภาพป้าย + เลขที่แก้แล้ว ไว้ใช้เทรน OCR รอบถัดไป (data/corrections/)."""
    img_dir = CORRECTIONS_DIR / "images"
    img_dir.mkdir(parents=True, exist_ok=True)
    rows = db.query(Event).filter(Event.is_corrected.is_(True), Event.plate_img.isnot(None)).all()
    with open(CORRECTIONS_DIR / "labels.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["image", "plate_text", "plate_province", "ocr_text", "ocr_province"])
        for e in rows:
            src = DATA_DIR / e.plate_img
            name = f"{e.id:06d}.jpg"
            if src.exists() and not (img_dir / name).exists():
                shutil.copy(src, img_dir / name)
            w.writerow([f"images/{name}", e.plate_text or "", e.plate_province or "", e.original_plate_text or "",
                        e.original_plate_province or ""])
